"""Import the supplied spreadsheet once. The three PII columns are never read (spec §3.2)."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from scout.domain.listing import (
    AreaBasis,
    BhkType,
    Coordinates,
    Furnishing,
    ListingRecord,
    Parking,
    PropertyType,
    SocietyType,
)
from scout.pipeline.pii import strip_pii


class SheetSchemaError(ValueError):
    """Raised when a required column is absent."""


SHEET = "Bangalore_Properties_List"
# Every column the record cannot be built honestly without. `availability_status` and
# `Society Type` joined the sheet on 2026-09-05: Gate D was re-decided on the strength of
# the availability marker, so if either column ever vanishes the import must fail rather
# than write a null that reads as "not stated" and quietly stops curation dropping rows.
REQUIRED = (
    "locality",
    "bhk_type",
    "Rent",
    "Deposit",
    "Latitude",
    "Longitude",
    "Society Type",
    "availability_status",
)
SOURCE_PATH = "data/Bangalore_Properties_List.xlsx"

# Columns the owner added on 2026-09-05 that hold personal data. They are named here so
# the allow-list below is checkable, and they are never read into a record.
PII_COLUMNS = frozenset({"Name", "Phone Number", "Voter ID"})

# The only columns the importer may read. The three PII columns above are absent on
# purpose. This is the first line of defence; strip_pii is the second.
IMPORTED_COLUMNS = frozenset(
    {
        "Sl.",
        "locality",
        "property_type",
        "bhk_type",
        "bedrooms",
        "bathrooms",
        "balconies",
        "square_feet",
        "Rent",
        "Deposit",
        "furnishing",
        "parking_available",
        "society_name",
        "Society Type",
        "total_floors",
        "Latitude",
        "Longitude",
        "availability_status",
    }
)

# Text a source uses to say "this field does not apply"; it is null, never a number.
_NOT_APPLICABLE = {"", "not applicable", "n/a", "na", "-"}


def _cell(cells: tuple[Any, ...], col: dict[str, int], name: str) -> Any:
    """Read one cell by column name. A name outside the allow-list is a programming error."""
    if name not in IMPORTED_COLUMNS:
        raise KeyError(f"{name!r} is not an imported column; see IMPORTED_COLUMNS")
    i = col.get(name)
    return cells[i] if i is not None and i < len(cells) else None


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def _text(v: Any) -> str | None:
    if v is None:
        return None
    s = strip_pii(str(v)).strip()
    return s or None


def _int(v: Any, *, row: int, col: str) -> int | None:
    if v is None:
        return None
    if isinstance(v, bool):
        raise TypeError(f"row {row}: {col} is a boolean, expected a number")
    if isinstance(v, (int, float)):
        if float(v) != int(v):
            raise ValueError(f"row {row}: {col} is not a whole number: {v!r}")
        return int(v)
    s = str(v).strip()
    if s.lower() in _NOT_APPLICABLE:
        return None
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    raise ValueError(f"row {row}: {col} is not a number: {s!r}")


def _float(v: Any, *, row: int, col: str) -> float | None:
    if v is None or (isinstance(v, str) and v.strip().lower() in _NOT_APPLICABLE):
        return None
    if isinstance(v, bool):
        raise TypeError(f"row {row}: {col} is a boolean, expected a number")
    try:
        return float(v)
    except (TypeError, ValueError) as e:
        raise ValueError(f"row {row}: {col} is not a number: {v!r}") from e


def _yes_no(v: Any, *, row: int, col: str) -> bool | None:
    s = _text(v)
    if s is None:
        return None
    if s.casefold() == "yes":
        return True
    if s.casefold() == "no":
        return False
    raise ValueError(f"row {row}: {col} must be Yes or No, got {s!r}")


# The sheet spells society type in prose. Only these two spellings exist in the source
# (4,590 rows each, no blanks, measured 2026-09-05); anything else is a change worth failing on.
_SOCIETY_TYPE = {
    "gated society": SocietyType.GATED,
    "non-gated society": SocietyType.NON_GATED,
}


def _society_type(v: Any, *, row: int, col: str) -> SocietyType | None:
    s = _text(v)
    if s is None:
        return None
    try:
        return _SOCIETY_TYPE[s.casefold()]
    except KeyError as e:
        allowed = ", ".join(sorted(_SOCIETY_TYPE))
        raise ValueError(f"row {row}: {col} must be one of {allowed}, got {s!r}") from e


def _enum(kind: type, v: Any, *, row: int, col: str) -> Any:
    s = _text(v)
    if s is None:
        return None
    try:
        return kind(s)
    except ValueError as e:
        allowed = ", ".join(m.value for m in kind)
        raise ValueError(f"row {row}: {col} must be one of {allowed}, got {s!r}") from e


def _record(cells: tuple[Any, ...], col: dict[str, int], as_of: date, row: int) -> ListingRecord:
    def get(name: str) -> Any:
        return _cell(cells, col, name)

    locality = _text(get("locality"))
    if locality is None:
        raise ValueError(f"row {row}: locality is blank")
    serial = _int(get("Sl."), row=row, col="Sl.")
    if serial is None:
        serial = row  # the Excel row number stands in when the sheet carries no serial
    lat = _float(get("Latitude"), row=row, col="Latitude")
    lng = _float(get("Longitude"), row=row, col="Longitude")
    coordinates = Coordinates(lat=lat, lng=lng) if lat is not None and lng is not None else None
    # The sheet says only that parking exists. The kind stays None: a bare "Yes" is
    # never inflated into Parking.BOTH.
    parking_kind: Parking | None = None

    return ListingRecord(
        id=f"{_slug(locality)}-{serial:05d}",
        source_url=f"file://{SOURCE_PATH}#row={serial}",
        scraped_on=as_of,
        locality=locality,
        bhk_type=_enum(BhkType, get("bhk_type"), row=row, col="bhk_type"),
        bedrooms=_int(get("bedrooms"), row=row, col="bedrooms"),
        bathrooms=_int(get("bathrooms"), row=row, col="bathrooms"),
        balconies=_int(get("balconies"), row=row, col="balconies"),
        rent=_int(get("Rent"), row=row, col="Rent"),
        deposit=_int(get("Deposit"), row=row, col="Deposit"),
        property_type=_enum(PropertyType, get("property_type"), row=row, col="property_type"),
        furnishing=_enum(Furnishing, get("furnishing"), row=row, col="furnishing"),
        square_footage=_int(get("square_feet"), row=row, col="square_feet"),
        area_basis=AreaBasis.UNKNOWN,
        total_floors=_int(get("total_floors"), row=row, col="total_floors"),
        parking=parking_kind,
        parking_available=_yes_no(get("parking_available"), row=row, col="parking_available"),
        availability_status=_yes_no(get("availability_status"), row=row, col="availability_status"),
        society_name=_text(get("society_name")),
        society_type=_society_type(get("Society Type"), row=row, col="Society Type"),
        coordinates=coordinates,
    )


def import_sheet(path: Path, as_of: date) -> list[ListingRecord]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise SheetSchemaError(f"{path}: no sheet named {SHEET!r} (found {wb.sheetnames})")
        ws = wb[SHEET]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise SheetSchemaError(f"{path}: the sheet is empty")
        col = {str(h).strip(): i for i, h in enumerate(header) if h is not None}
        missing = [name for name in REQUIRED if name not in col]
        if missing:
            raise SheetSchemaError(f"{path}: missing required column(s): {', '.join(missing)}")
        records: list[ListingRecord] = []
        for row, cells in enumerate(rows, start=2):  # row 1 is the header
            if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
                continue
            records.append(_record(cells, col, as_of, row))
        return records
    finally:
        wb.close()


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", default=SOURCE_PATH)
    ap.add_argument("--out", default="data/raw/listings_all.json")
    ap.add_argument(
        "--as-of",
        default=datetime.now(ZoneInfo("Asia/Kolkata")).date().isoformat(),
        help="the date the sheet is taken as of (YYYY-MM-DD); defaults to today in IST",
    )
    args = ap.parse_args()
    recs = import_sheet(Path(args.sheet), date.fromisoformat(args.as_of))
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(
        json.dumps([r.model_dump(mode="json") for r in recs], ensure_ascii=False, indent=1),
        encoding="utf-8",
    )
    localities = {r.locality for r in recs}
    print(f"imported {len(recs)} records over {len(localities)} localities -> {out}")
