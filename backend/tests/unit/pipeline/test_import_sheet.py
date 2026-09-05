from datetime import date
from pathlib import Path

import pytest

from scout.domain.listing import AreaBasis, BhkType
from scout.pipeline.import_sheet import SheetSchemaError, import_sheet

SAMPLE = Path(__file__).parents[2] / "fixtures" / "listing_sample.xlsx"


def test_maps_every_supplied_column():
    recs = import_sheet(SAMPLE, date(2026, 9, 2))
    assert len(recs) == 5
    r = recs[0]
    assert isinstance(r.locality, str) and r.locality
    assert r.rent is not None
    assert r.deposit is not None
    assert r.bhk_type in set(BhkType)
    assert r.coordinates is not None


def test_fields_the_sheet_lacks_are_none_not_guessed():
    r = import_sheet(SAMPLE, date(2026, 9, 2))[0]
    assert r.amenities is None
    assert r.lift is None
    assert r.floor is None
    assert r.available_from is None
    assert r.maintenance_charges is None
    assert r.area_basis is AreaBasis.UNKNOWN


def test_a_bare_yes_never_becomes_a_four_wheeler_claim():
    for r in import_sheet(SAMPLE, date(2026, 9, 2)):
        assert r.parking is None
        assert r.parking_available in {True, False}


def test_balconies_is_imported():
    for r in import_sheet(SAMPLE, date(2026, 9, 2)):
        assert isinstance(r.balconies, int)


def test_output_carries_no_pii():
    dumped = "".join(r.model_dump_json() for r in import_sheet(SAMPLE, date(2026, 9, 2)))
    assert "@" not in dumped
    import re

    # no 10-digit mobiles anywhere
    assert not re.search(r"(?<!\d)[6-9]\d{9}(?!\d)", dumped)


def test_missing_required_column_is_an_error(tmp_path):
    from openpyxl import Workbook, load_workbook

    src = load_workbook(SAMPLE, read_only=True)[
        "Bangalore_Properties_List"
    ]
    rows = list(src.iter_rows(values_only=True))
    header = list(rows[0])
    drop = header.index("locality")
    wb = Workbook()
    ws = wb.active
    ws.title = "Bangalore_Properties_List"
    for row in rows:
        ws.append([c for i, c in enumerate(row) if i != drop])
    broken = tmp_path / "no_locality.xlsx"
    wb.save(broken)

    with pytest.raises(SheetSchemaError):
        import_sheet(broken, date(2026, 9, 2))


def test_availability_status_is_imported():
    recs = import_sheet(SAMPLE, date(2026, 9, 2))
    assert [r.availability_status for r in recs] == [True, True, True, False, False]


def test_pii_columns_are_never_imported():
    dumped = "".join(r.model_dump_json() for r in import_sheet(SAMPLE, date(2026, 9, 2)))
    for value in ("Zoya Narang", "Meera Nair", "MPSAnXINZQ", "395862397", "504726183"):
        assert value not in dumped


def test_the_allow_list_excludes_every_pii_column():
    from scout.pipeline.import_sheet import IMPORTED_COLUMNS, PII_COLUMNS

    assert PII_COLUMNS == {"Name", "Phone Number", "Voter ID"}
    assert IMPORTED_COLUMNS.isdisjoint(PII_COLUMNS)


def test_reading_a_column_outside_the_allow_list_is_an_error():
    from scout.pipeline.import_sheet import _cell

    with pytest.raises(KeyError):
        _cell(("x",), {"Phone Number": 0}, "Phone Number")


def test_society_type_is_mapped_not_stored_verbatim():
    from scout.domain.listing import SocietyType

    recs = import_sheet(SAMPLE, date(2026, 9, 2))
    assert [r.society_type for r in recs] == [
        SocietyType.GATED,
        SocietyType.NON_GATED,
        SocietyType.GATED,
        SocietyType.NON_GATED,
        SocietyType.GATED,
    ]


def test_an_unrecognised_society_type_is_an_error(tmp_path):
    from openpyxl import Workbook, load_workbook

    src = load_workbook(SAMPLE, read_only=True)["Bangalore_Properties_List"]
    rows = [list(r) for r in src.iter_rows(values_only=True)]
    i = [str(h).strip() for h in rows[0]].index("Society Type")
    rows[1][i] = "Cooperative Society"
    wb = Workbook()
    ws = wb.active
    ws.title = "Bangalore_Properties_List"
    for r in rows:
        ws.append(r)
    p = tmp_path / "bad_society_type.xlsx"
    wb.save(p)
    with pytest.raises(ValueError, match="Society Type must be one of"):
        import_sheet(p, date(2026, 9, 2))


def _sheet_without(tmp_path, column):
    """Copy the fixture with one column removed."""
    from openpyxl import Workbook, load_workbook

    src = load_workbook(SAMPLE, read_only=True)["Bangalore_Properties_List"]
    rows = list(src.iter_rows(values_only=True))
    drop = [str(h).strip() for h in rows[0]].index(column)
    wb = Workbook()
    ws = wb.active
    ws.title = "Bangalore_Properties_List"
    for row in rows:
        ws.append([c for i, c in enumerate(row) if i != drop])
    p = tmp_path / "missing.xlsx"
    wb.save(p)
    return p


@pytest.mark.parametrize("column", ["availability_status", "Society Type"])
def test_a_missing_2026_09_05_column_is_an_error_not_a_silent_null(tmp_path, column):
    # Gate D was re-decided on the strength of `availability_status`, and `society_type`
    # was added to the schema because the sheet carries `Society Type`. If either column
    # ever disappears the importer must fail loudly: a silent null availability would
    # quietly stop curation dropping unavailable rows (spec §3.1) and nothing would say so.
    with pytest.raises(SheetSchemaError, match=column):
        import_sheet(_sheet_without(tmp_path, column), date(2026, 9, 2))
